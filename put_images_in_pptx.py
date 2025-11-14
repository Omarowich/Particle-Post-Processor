import re, math
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# ── CONFIG ─────────────────────────────────────────────────────────────
images_root = Path(r"\\nas.ads.mwn.de\tuei\mml\MML MS BS students\Bachelor Students\Omar Elsabbagh\HIWI\NAF\normalised analysis")  # your folder
output_pptx = images_root / "plots_by_type_Tkb.pptx"

cols     = 4
margin   = Inches(0.2)
title_h  = Inches(0.5)
label_h  = Inches(0.6)   # reserve this much height _above_ each plot for the Tkb label
row_gap  = Inches(0.4)   # space _below_ each plot before the next row
# ────────────────────────────────────────────────────────────────────────

def parse_tkb(fn):
    m = re.match(r"([\d\.]+)Tkb", fn)
    return float(m.group(1)) if m else float("inf")

# 1) collect & sort your Tkb folders
param_dirs = sorted([d for d in images_root.iterdir() if d.is_dir()],
                    key=lambda d: parse_tkb(d.name))

# 2) collect all plot‐stems
plot_stems = sorted({img.stem
                     for d in param_dirs
                     for img in d.glob("*.png")})

# 3) compute grid & slide height
n_params = len(param_dirs)
n_rows   = math.ceil(n_params/cols)
prs      = Presentation()

usable_w = prs.slide_width - 2*margin
img_w    = (usable_w - (cols-1)*margin)/cols
img_h    = img_w

# total height = top margin + title + gap + rows*(label + plot + gap) + bottom margin
total_h = (margin
           + title_h
           + margin
           + n_rows*(label_h + img_h + row_gap)
           + margin)
prs.slide_height = int(round(total_h))

blank = prs.slide_layouts[6]

# 4) build slides
for stem in plot_stems:
    slide = prs.slides.add_slide(blank)

    # slide title
    tb = slide.shapes.add_textbox(margin, margin/2,
                                  prs.slide_width-2*margin, title_h)
    p  = tb.text_frame.add_paragraph()
    p.text      = stem.replace("_", " ").capitalize()
    p.font.size = Pt(28)
    p.alignment = PP_ALIGN.CENTER

    # place each Tkb in ascending order
    for idx, d in enumerate(param_dirs):
        img_path = d / f"{stem}.png"
        if not img_path.exists():
            continue

        row, col = divmod(idx, cols)
        left     = margin + col*(img_w + margin)
        top_base = margin + title_h + margin + row*(label_h + img_h + row_gap)

        # 4a) label ABOVE the plot
        lbl_tb = slide.shapes.add_textbox(left, top_base, img_w, label_h)
        lbl_p  = lbl_tb.text_frame.add_paragraph()
        lbl_p.text      = f"{parse_tkb(d.name)} Tkb"
        lbl_p.font.size = Pt(12)
        lbl_p.alignment = PP_ALIGN.LEFT

        # 4b) picture immediately _below_ that label
        slide.shapes.add_picture(
            str(img_path),
            left,
            top_base + label_h,
            width=img_w
        )

prs.save(output_pptx)
print(f"✅ Saved: {output_pptx}")
