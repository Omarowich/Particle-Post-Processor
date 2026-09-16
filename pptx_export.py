"""Build a PowerPoint deck from a folder of per-run plot PNGs (one slide per
metric, one column per Tkb folder).

Extracted from app.py (Phase 1 de-spaghetti pass).
"""

import math
import re
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def build_pptx_from_images(
    images_root: Path,
    output_pptx: Optional[Path] = None,
    cols: int = 4,
    margin_in=0.2,
    title_h_in=0.5,
    label_h_in=0.6,
    row_gap_in=0.4,
) -> Path:
    """
    Walk a folder structure like:
        images_root/
            0Tkb_.../
                metric1.png
                metric2.png
            5Tkb_.../
                metric1.png
                metric2.png
            ...

    and build a PPTX where each slide is one 'metric' (stem name),
    with columns = different Tkb folders.
    """

    margin  = Inches(margin_in)
    title_h = Inches(title_h_in)
    label_h = Inches(label_h_in)
    row_gap = Inches(row_gap_in)

    def parse_tkb(fn: str) -> float:
        m = re.match(r"([\d\.]+)Tkb", fn)
        return float(m.group(1)) if m else float("inf")

    if output_pptx is None:
        output_pptx = images_root / "plots_by_type_Tkb.pptx"

    # 1) collect & sort your Tkb folders
    param_dirs = sorted(
        [d for d in images_root.iterdir() if d.is_dir()],
        key=lambda d: parse_tkb(d.name),
    )
    if not param_dirs:
        raise RuntimeError(f"No Tkb folders found inside: {images_root}")

    # 2) collect all plot‐stems
    plot_stems = sorted(
        {
            img.stem
            for d in param_dirs
            for img in d.glob("*.png")
        }
    )
    if not plot_stems:
        raise RuntimeError(f"No .png images found under: {images_root}")

    n_params = len(param_dirs)
    n_rows   = math.ceil(n_params / cols)

    prs = Presentation()

    usable_w = prs.slide_width - 2*margin
    img_w    = (usable_w - (cols-1)*margin) / cols
    img_h    = img_w

    # total height = top margin + title + gap + rows*(label + plot + gap) + bottom margin
    total_h = (
        margin
        + title_h
        + margin
        + n_rows * (label_h + img_h + row_gap)
        + margin
    )
    prs.slide_height = int(round(total_h))

    blank = prs.slide_layouts[6]

    # 4) build slides
    for stem in plot_stems:
        slide = prs.slides.add_slide(blank)

        # slide title
        tb = slide.shapes.add_textbox(
            margin,
            margin / 2,
            prs.slide_width - 2 * margin,
            title_h,
        )
        p  = tb.text_frame.add_paragraph()
        p.text      = stem.replace("_", " ").capitalize()
        p.font.size = Pt(28)
        p.alignment = PP_ALIGN.CENTER

        # place each Tkb in ascending order
        for idx, d in enumerate(param_dirs):
            img_path = d / f"{stem}.png"
            if not img_path.exists():
                continue

            row, col = divmod(idx, cols)
            left     = margin + col * (img_w + margin)
            top_base = margin + title_h + margin + row * (label_h + img_h + row_gap)

            # label ABOVE the plot
            lbl_tb = slide.shapes.add_textbox(left, top_base, img_w, label_h)
            lbl_p  = lbl_tb.text_frame.add_paragraph()
            lbl_p.text      = f"{parse_tkb(d.name)} Tkb"
            lbl_p.font.size = Pt(12)
            lbl_p.alignment = PP_ALIGN.LEFT

            # picture immediately below that label
            slide.shapes.add_picture(
                str(img_path),
                left,
                top_base + label_h,
                width=img_w,
            )

    prs.save(output_pptx)
    return output_pptx
